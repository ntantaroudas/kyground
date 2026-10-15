"""Optional manual review workbook (export) and import of a reviewer's decisions.

The gold answers are verified automatically against the frozen snapshot (``bench/verify.py``); this workbook is
for any additional manual review (for example by a Greek-speaking colleague) and is not part of the study's
protocol. The workbook is in Greek; editable columns are highlighted in yellow and the decision column has a
drop-down list. A summary sheet counts decisions with live formulas. An item marked wrong or unclear stays
unverified however the automatic checks turn out.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..schema import Item, load_items, write_jsonl
from .nuggets import describe

DECISIONS = ("Σωστό", "Λάθος", "Ασαφές")
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")

COLUMNS = [
    # (header, width, editable)
    ("Κωδικός (iid)", 34, False),
    ("Τύπος", 14, False),
    ("Εγγραφές (πηγή)", 34, False),
    ("Σύνδεσμος", 34, False),
    ("Ερώτηση (ΕΛ)", 54, False),
    ("Question (EN)", 46, False),
    ("Σωστή απάντηση (ΕΛ)", 46, False),
    ("Τι πρέπει να περιέχει η απάντηση", 46, False),
    ("Απόσπασμα από την πλατφόρμα", 46, False),
    ("ΕΛΕΓΧΟΣ", 12, True),
    ("Διορθωμένη απάντηση (αν χρειάζεται)", 40, True),
    ("Πιο φυσική διατύπωση της ερώτησης (προαιρετικό)", 40, True),
    ("Τοπική / ιδιωματική διατύπωση (προαιρετικό)", 40, True),
    ("Σχόλια", 30, True),
]


def _row(it: Item) -> list[str]:
    return [
        it.iid, it.qtype, ", ".join(it.gold_rids) or "—", it.source_url,
        it.question.get("el", ""), it.question.get("en", ""), it.answer_ref.get("el", ""),
        " · ".join(describe(n, "el") for n in it.nuggets), it.evidence.get("el", ""),
        "", "", "", "", "",
    ]


def export_review(items: list[Item], path: str | Path, title: str = "KyGround – έλεγχος απαντήσεων") -> Path:
    wb = Workbook()
    # ---------------------------------------------------------- instructions
    ins = wb.active
    ins.title = "Οδηγίες"
    lines = [
        (title, True),
        ("", False),
        ("Τι ζητάμε: για κάθε γραμμή του φύλλου «Ερωτήσεις», ελέγξτε αν η «Σωστή απάντηση» είναι σωστή σύμφωνα "
         "με το περιεχόμενο της πλατφόρμας DASK.", False),
        ("Συμπληρώνετε ΜΟΝΟ τις κίτρινες στήλες (J–N). Οι υπόλοιπες στήλες δεν αλλάζουν.", False),
        ("Στήλη J «ΕΛΕΓΧΟΣ»: επιλέξτε από τη λίστα Σωστό, Λάθος ή Ασαφές.", False),
        ("Στήλη K: αν η απάντηση είναι λάθος ή ελλιπής, γράψτε τη σωστή.", False),
        ("Στήλη L (προαιρετικό): πώς θα έκανε την ερώτηση ένας παραγωγός ή επισκέπτης στην καθημερινή γλώσσα.", False),
        ("Στήλη M (προαιρετικό): η ίδια ερώτηση με τοπικές/κυθηραϊκές λέξεις, αν υπάρχουν.", False),
        ("Στήλη N: οποιοδήποτε σχόλιο (π.χ. η εγγραφή είναι δοκιμαστική και πρέπει να αφαιρεθεί).", False),
        ("Οι ερωτήσεις τύπου «unanswerable» αφορούν πληροφορίες που ΔΕΝ υπάρχουν στην πλατφόρμα: "
         "επιβεβαιώστε ότι πράγματι δεν υπάρχουν.", False),
        ("Παράδειγμα συμπληρωμένης γραμμής: J = Σωστό · K = (κενό) · L = «Πότε μαζεύουμε τις ελιές;» · "
         "M = «Πότε είναι το λιομάζωμα στο Τσιρίγο;» · N = (κενό)", False),
        ("Το φύλλο «Σύνοψη» μετρά αυτόματα πόσες γραμμές έχουν ελεγχθεί.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=txt)
        c.font = Font(name=FONT, size=13 if bold else 11, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ins.column_dimensions["A"].width = 120

    # ----------------------------------------------------------------- items
    ws = wb.create_sheet("Ερωτήσεις")
    for j, (h, w, editable) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FONT, bold=True)
        c.fill = YELLOW if editable else HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.cell(row=1, column=10).comment = Comment("Επιλέξτε Σωστό / Λάθος / Ασαφές", "KyGround")
    for i, it in enumerate(items, start=2):
        for j, v in enumerate(_row(it), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            c.fill = YELLOW if COLUMNS[j - 1][2] else (GREY if j <= 4 else PatternFill())
    last = len(items) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dv.error = "Επιλέξτε από τη λίστα"
    dv.prompt = "Σωστό / Λάθος / Ασαφές"
    ws.add_data_validation(dv)
    dv.add(f"J2:J{last}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:N{last}"

    # --------------------------------------------------------------- summary
    sm = wb.create_sheet("Σύνοψη")
    rows = [
        ("Σύνολο ερωτήσεων", f"=COUNTA('Ερωτήσεις'!A2:A{last})"),
        ("Σωστό", f"=COUNTIF('Ερωτήσεις'!J2:J{last},\"Σωστό\")"),
        ("Λάθος", f"=COUNTIF('Ερωτήσεις'!J2:J{last},\"Λάθος\")"),
        ("Ασαφές", f"=COUNTIF('Ερωτήσεις'!J2:J{last},\"Ασαφές\")"),
        ("Δεν έχουν ελεγχθεί", "=B2-B3-B4-B5"),
        ("Ποσοστό ελεγμένων", "=IF(B2=0,0,(B3+B4+B5)/B2)"),
    ]
    for i, (k, f) in enumerate(rows, start=2):
        a = sm.cell(row=i, column=1, value=k)
        b = sm.cell(row=i, column=2, value=f)
        a.font = Font(name=FONT, bold=(i == 2))
        b.font = Font(name=FONT)
        if k.startswith("Ποσοστό"):
            b.number_format = "0.0%"
    sm.cell(row=1, column=1, value="Μέτρηση").font = Font(name=FONT, bold=True)
    sm.cell(row=1, column=2, value="Τιμή").font = Font(name=FONT, bold=True)
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 14
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_review(items_path: str | Path, review_path: str | Path, out_items: str | Path,
                  decisions_path: str | Path | None = None) -> dict:
    """Apply a reviewer's decisions to items.jsonl and record them in ``review_decisions.yaml`` so that they
    survive a rebuild of the benchmark (``kyground build-bench`` re-applies them).

    Σωστό -> verified; Λάθος -> needs_fix (with a correction) or rejected (removed at the next build);
    Ασαφές -> unclear. Corrections are not applied automatically: fix the template or the hand-written item and
    rebuild. Only verified items are used in real runs (``accuracy.item_statuses`` in configs/study.yaml), and
    every rebuild re-runs the automatic verification.
    """
    import datetime
    import yaml
    items = {it.iid: it for it in load_items(items_path)}
    wb = load_workbook(review_path, data_only=True)
    ws = wb["Ερωτήσεις"]
    counts = {"verified": 0, "rejected": 0, "needs_fix": 0, "unclear": 0, "unchanged": 0, "corrected": 0}
    decisions: dict[str, dict] = {}
    if decisions_path and Path(decisions_path).exists():
        decisions = yaml.safe_load(Path(decisions_path).read_text(encoding="utf-8")) or {}
    today = datetime.date.today().isoformat()
    for row in ws.iter_rows(min_row=2, values_only=True):
        iid = row[0]
        if not iid or iid not in items:
            continue
        it = items[iid]
        decision, corr, natural, local, comment = (row[9] or "", row[10] or "", row[11] or "", row[12] or "",
                                                   row[13] or "")
        if decision == "Σωστό":
            it.status = "verified"
        elif decision == "Λάθος":
            it.status = "needs_fix" if corr else "rejected"
        elif decision == "Ασαφές":
            it.status = "unclear"
        else:
            counts["unchanged"] += 1
            continue
        counts[it.status] += 1
        if corr:
            it.notes = (it.notes + " | " if it.notes else "") + f"reviewer correction: {corr}"
            counts["corrected"] += 1
        if natural:
            it.notes = (it.notes + " | " if it.notes else "") + f"original question: {it.question['el']}"
            it.question["el"] = str(natural).strip()
        if local:
            it.variants["el_colloq"] = str(local).strip()
        if comment:
            it.notes = (it.notes + " | " if it.notes else "") + f"reviewer: {comment}"
        decisions[iid] = {k: v for k, v in {"status": it.status, "correction": str(corr).strip(),
                                            "natural_question": str(natural).strip(),
                                            "local_variant": str(local).strip(), "comment": str(comment).strip(),
                                            "date": today}.items() if v}
    write_jsonl(out_items, items.values())
    if decisions_path:
        Path(decisions_path).write_text(yaml.safe_dump(decisions, allow_unicode=True, sort_keys=True),
                                        encoding="utf-8")
    return counts


def apply_decisions(items: list, decisions: dict) -> list:
    """Re-apply recorded review decisions after a rebuild; rejected items are removed."""
    out = []
    for it in items:
        d = decisions.get(it.iid)
        if d:
            if d.get("status") == "rejected":
                continue
            it.status = d.get("status", it.status)
            if d.get("natural_question"):
                it.notes = (it.notes + " | " if it.notes else "") + f"original question: {it.question['el']}"
                it.question["el"] = d["natural_question"]
            if d.get("local_variant"):
                it.variants["el_colloq"] = d["local_variant"]
            for k, lab in (("correction", "reviewer correction"), ("comment", "reviewer")):
                if d.get(k):
                    it.notes = (it.notes + " | " if it.notes else "") + f"{lab}: {d[k]}"
        out.append(it)
    return out
